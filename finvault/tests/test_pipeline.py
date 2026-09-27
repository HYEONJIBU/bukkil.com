import pytest

from finvault import db, ingest
from finvault.analysis import Vault
from finvault.client import DartClient, DartError, RateLimitError

from fake_dart import FakeResponse, FakeSession, revenue


@pytest.fixture
def env(tmp_path):
    path = tmp_path / "t.db"
    conn = db.connect(path)
    session = FakeSession()
    client = DartClient("KEY", min_interval=0, session=session)
    ingest.ingest_company(conn, client, "테스트전자", 2023, 2024, log=lambda *_: None)
    return conn, session, Vault(path)


def test_resolve_prefers_exact_listed(env):
    conn, _, _ = env
    assert ingest.resolve_corp(conn, "테스트전자")["corp_code"] == "00000001"
    assert ingest.resolve_corp(conn, "999990")["corp_code"] == "00000001"
    assert ingest.resolve_corp(conn, "테스트전자서비스")["corp_code"] == "00000002"


def test_quarterly_and_annual_revenue(env):
    _, _, vault = env
    q = vault.metrics("테스트전자", ["revenue", "operating_income"], freq="Q")
    assert list(q.index) == [f"{y}Q{i}" for y in (2023, 2024) for i in (1, 2, 3, 4)]
    for y in (2023, 2024):
        for i in (1, 2, 3, 4):
            assert q.loc[f"{y}Q{i}", "revenue"] == pytest.approx(revenue(y, i))
            assert q.loc[f"{y}Q{i}", "operating_income"] == pytest.approx(revenue(y, i) * 0.1)
    a = vault.metrics("테스트전자", ["revenue"], freq="A")
    assert a.loc["2024A", "revenue"] == pytest.approx(sum(revenue(2024, i) for i in (1, 2, 3, 4)))
    assert a.attrs["fs_div"] == "CFS"


def test_non_standard_net_income_joined_across_quarters(env):
    _, _, vault = env
    q = vault.metrics("테스트전자", ["net_income"], freq="Q")
    assert q.loc["2024Q4", "net_income"] == pytest.approx(revenue(2024, 4) * 0.08)


def test_cash_flow_quarterly(env):
    _, _, vault = env
    q = vault.metrics("테스트전자", ["cfo"], freq="Q")
    assert q.loc["2024Q3", "cfo"] == pytest.approx(revenue(2024, 3) * 0.2)


def test_separate_statements_and_units(env):
    _, _, vault = env
    a = vault.metrics("테스트전자", ["revenue"], freq="A", fs_div="OFS", unit="백만")
    assert a.loc["2023A", "revenue"] == pytest.approx(sum(revenue(2023, i) for i in (1, 2, 3, 4)) * 0.5 / 1e6)


def test_ratios_and_ttm(env):
    _, _, vault = env
    r = vault.ratios("테스트전자", freq="A")
    assert r.loc["2024A", "operating_margin"] == pytest.approx(10)
    assert r.loc["2024A", "debt_ratio"] == pytest.approx(4000 / 6040 * 100)
    rq = vault.ratios("테스트전자", freq="Q")
    assert rq.loc["2024Q2", "revenue_yoy"] == pytest.approx((revenue(2024, 2) / revenue(2023, 2) - 1) * 100)
    t = vault.ttm("테스트전자", ["revenue"])
    assert t.loc["2024Q2", "revenue"] == pytest.approx(
        revenue(2023, 3) + revenue(2023, 4) + revenue(2024, 1) + revenue(2024, 2))


def test_custom_account_and_items(env):
    _, _, vault = env
    acc = vault.account("테스트전자", "연구개발", freq="Q")
    assert acc.iloc[-1, 0] == pytest.approx(revenue(2024, 4) * 0.05, abs=1)  # 원 단위 반올림
    div = vault.items("테스트전자", "alotMatter")
    assert set(div["bsns_year"]) == {2023, 2024}
    assert vault.sql("SELECT COUNT(*) AS n FROM disclosures")["n"][0] == 1


def test_refetch_skipped_unless_refresh(env):
    conn, session, _ = env
    n = len(session.calls)
    client = DartClient("KEY", min_interval=0, session=session)
    ingest.ingest_company(conn, client, "테스트전자", 2023, 2024, log=lambda *_: None)
    fs_calls = [c for c in session.calls[n:] if c[0] == "fnlttSinglAcntAll.json"]
    assert fs_calls == []


def test_client_status_handling():
    class S:
        def __init__(self, payload):
            self.payload = payload

        def get(self, *a, **k):
            return FakeResponse(self.payload)

    assert DartClient("K", 0, session=S({"status": "013", "message": ""})).get_list("x") == []
    with pytest.raises(RateLimitError):
        DartClient("K", 0, session=S({"status": "020", "message": "limit"})).get_list("x")
    with pytest.raises(DartError):
        DartClient("K", 0, session=S({"status": "010", "message": "bad key"})).get_list("x")


def test_excel_export(env, tmp_path):
    import openpyxl
    from finvault.export import export_company

    _, _, vault = env
    path = export_company("테스트전자", tmp_path / "out.xlsx", vault=vault)
    wb = openpyxl.load_workbook(path)
    assert {"안내", "요약_연간", "요약_분기", "손익계산서_분기", "재무상태표_연간", "배당에 관한 사항"} <= set(wb.sheetnames)
    rows = {r[0]: r[1:] for r in wb["요약_연간"].iter_rows(min_row=2, values_only=True)}
    assert rows["매출액(억원)"][0] == pytest.approx(sum(revenue(2023, i) for i in (1, 2, 3, 4)) / 1e8)
